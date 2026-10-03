"""
Moderne oppsett i tilbudsmalen: bort med rutenettet, inn med luft og hierarki.

Prinsippet er at linjer skal skille der oyet trenger hjelp, og ellers vaere
borte. Et rutenett rundt hver celle gjor det motsatte - det rammer inn hvert
tall som om det var like viktig som alle andre.
"""
import copy
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Mm

MAL = 'maler/tilbud-mal.docx'

MORK   = RGBColor(0x12, 0x10, 0x0C)   # BWW-svart, fra logoen
GRA    = RGBColor(0x76, 0x76, 0x76)   # etiketter og sekundaertekst
HARLIN = 'D9D9D9'                     # radskille
STREK  = '12100C'                     # hovedstrek under kolonneoverskriftene
GUL    = 'FFCF0D'                     # BWW-gul, brukt EN gang


def _el(tag, **attr):
    e = OxmlElement(tag)
    for k, v in attr.items():
        e.set(qn('w:' + k), v)
    return e


def kant(sider):
    """sider: {'bottom': ('single', 4, 'D9D9D9'), 'top': ('nil',)}"""
    b = OxmlElement('w:tblBorders')
    for navn in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        spec = sider.get(navn, ('nil',))
        if spec[0] == 'nil':
            b.append(_el('w:' + navn, val='nil'))
        else:
            b.append(_el('w:' + navn, val=spec[0], sz=str(spec[1]),
                         space='0', color=spec[2]))
    return b


def sett_tabellkant(tabell, sider):
    pr = tabell._tbl.tblPr
    for tag in ('w:tblStyle', 'w:tblBorders'):
        for gammel in pr.findall(qn(tag)):
            pr.remove(gammel)
    pr.append(kant(sider))


def cellemarg(tabell, topp=0, bunn=0, venstre=0, hoyre=0):
    """Luft inne i cellene, i twips (1 mm = 56,7)."""
    pr = tabell._tbl.tblPr
    for gammel in pr.findall(qn('w:tblCellMar')):
        pr.remove(gammel)
    m = OxmlElement('w:tblCellMar')
    for navn, v in (('top', topp), ('bottom', bunn), ('left', venstre), ('right', hoyre)):
        m.append(_el('w:' + navn, w=str(v), type='dxa'))
    pr.append(m)


def radkant(rad, bunn=None, topp=None):
    """Strek under eller over EN rad, ved aa sette den paa hver celle."""
    for celle in rad.cells:
        pr = celle._tc.get_or_add_tcPr()
        for gammel in pr.findall(qn('w:tcBorders')):
            pr.remove(gammel)
        b = OxmlElement('w:tcBorders')
        if topp:
            b.append(_el('w:top', val='single', sz=str(topp[0]), space='0', color=topp[1]))
        if bunn:
            b.append(_el('w:bottom', val='single', sz=str(bunn[0]), space='0', color=bunn[1]))
        pr.append(b)


def tekst(run, storrelse=None, farge=None, fet=None, sperring=None, caps=False):
    if storrelse: run.font.size = Pt(storrelse)
    if farge is not None: run.font.color.rgb = farge
    if fet is not None: run.bold = fet
    if caps:
        run.text = run.text.upper()
    if sperring:
        rPr = run._element.get_or_add_rPr()
        for gammel in rPr.findall(qn('w:spacing')):
            rPr.remove(gammel)
        rPr.append(_el('w:spacing', val=str(sperring)))


def avstand(p, for_=None, etter=None):
    if for_ is not None: p.paragraph_format.space_before = Pt(for_)
    if etter is not None: p.paragraph_format.space_after = Pt(etter)


d = docx.Document(MAL)
kunde, varer, sum_t = d.tables[0], d.tables[1], d.tables[2]

# ---------------------------------------------------------------- topp
# Gul harstrek under logoen gir dokumentet en avsender uten aa rope.
logo_p = d.paragraphs[0]
avstand(logo_p, etter=6)
pPr = logo_p._p.get_or_add_pPr()
for gammel in pPr.findall(qn('w:pBdr')):      # idempotent: skriptet skal kunne
    pPr.remove(gammel)                        # kjores flere ganger uten aa hope seg
b = OxmlElement('w:pBdr')
b.append(_el('w:bottom', val='single', sz='12', space='4', color=GUL))
pPr.append(b)

tittel = d.paragraphs[1]
for r in tittel.runs:
    tekst(r, storrelse=26, farge=MORK, fet=True, sperring=40, caps=True)
avstand(tittel, for_=18, etter=2)

meta = d.paragraphs[2]
for r in meta.runs:
    tekst(r, storrelse=9, farge=GRA, fet=False)
avstand(meta, etter=14)

# ---------------------------------------------------------------- kunde
sett_tabellkant(kunde, {})
cellemarg(kunde, topp=60, bunn=60, venstre=0, hoyre=160)
for celle in kunde.rows[0].cells:
    for p in celle.paragraphs:
        for r in p.runs:
            tekst(r, storrelse=8, farge=GRA, fet=True, sperring=30, caps=True)
        avstand(p, etter=2)
radkant(kunde.rows[0], bunn=(4, HARLIN))
for celle in kunde.rows[1].cells:
    for p in celle.paragraphs:
        avstand(p, etter=0)
        for r in p.runs:
            tekst(r, storrelse=10, farge=MORK)

# ---------------------------------------------------------------- varer
sett_tabellkant(varer, {})
cellemarg(varer, topp=90, bunn=90, venstre=0, hoyre=140)
for celle in varer.rows[0].cells:
    for p in celle.paragraphs:
        for r in p.runs:
            tekst(r, storrelse=8, farge=GRA, fet=True, sperring=30, caps=True)
        avstand(p, etter=3)
radkant(varer.rows[0], bunn=(8, STREK))
# Datarada (indeks 2) er den docxtpl kopierer per vare - harlinja folger med.
radkant(varer.rows[2], bunn=(4, HARLIN))
for celle in varer.rows[2].cells:
    for p in celle.paragraphs:
        avstand(p, etter=0)
        for r in p.runs:
            tekst(r, storrelse=9.5, farge=MORK)

# ---------------------------------------------------------------- sum
sett_tabellkant(sum_t, {})
cellemarg(sum_t, topp=50, bunn=50, venstre=0, hoyre=140)
for i, rad in enumerate(sum_t.rows):
    siste = i == len(sum_t.rows) - 1
    for celle in rad.cells:
        for p in celle.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            avstand(p, etter=0)
            for r in p.runs:
                tekst(r, storrelse=12 if siste else 9.5,
                      farge=MORK if siste else GRA, fet=siste)
radkant(sum_t.rows[-1], topp=(8, STREK))

# ---------------------------------------------------------------- bunn
for r in d.paragraphs[-1].runs:
    tekst(r, storrelse=8, farge=GRA)
avstand(d.paragraphs[-1], for_=18)

d.save(MAL)
print("malen er omformet")


# ============================================================================
# Runde 2, etter forste rendering (LibreOffice -> PDF -> PNG).
#
# Tre feil som bare synes naar dokumentet faktisk tegnes opp:
#   1. Alle seks kolonnene var 29 mm. "BESKRIVELSE" brakk midt i ordet, mens
#      Varenr. hadde plass til overs.
#   2. Tallkolonnene var venstrestilt. I en pristabell skal tall staa paa
#      hoyre kant, ellers kan de ikke sammenlignes nedover.
#   3. Smaaskriften nederst falt ned paa en ellers tom side 2.
# ============================================================================
from docx.shared import Mm as _Mm

BREDDER_MM = [24, 52, 20, 24, 26, 28]        # sum 174 = tekstbredden
# Antall fikk 20 mm og ikke 16: med sperring i overskriftene brakk 'ANTALL'.
HOYRESTILT = [2, 3, 4, 5]                     # Antall, Veil., Din pris, Sum

d = docx.Document(MAL)
varer = d.tables[1]

# Kolonnebredder maa settes bade i tblGrid og paa hver enkelt celle - Word
# bruker cellene, LibreOffice ser paa griden. Settes bare den ene, spriker de.
grid = varer._tbl.find(qn('w:tblGrid'))
for kolonne, mm in zip(grid.findall(qn('w:gridCol')), BREDDER_MM):
    kolonne.set(qn('w:w'), str(int(mm * 56.7)))
varer.autofit = False
for rad in varer.rows:
    for celle, mm in zip(rad.cells, BREDDER_MM):
        celle.width = _Mm(mm)

for rad in varer.rows:
    for i in HOYRESTILT:
        for p in rad.cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

# Litt strammere rader, saa tabellen og smaaskriften faar plass paa side 1.
cellemarg(varer, topp=70, bunn=70, venstre=0, hoyre=110)
avstand(d.paragraphs[-1], for_=10)

d.save(MAL)
print("runde 2: kolonnebredder, hoyrestilte tall, strammere rader")
